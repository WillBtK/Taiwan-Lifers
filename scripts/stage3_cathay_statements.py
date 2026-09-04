#!/usr/bin/env python3
"""Stage 3 statements: Cathay Life quarterly Excel companions -> statement channel.

Cathay Holdings publishes Excel companions of Cathay Life's (5846) consolidated
quarterly statements from 2020 on (2012-2019 are PDF-only). Each workbook is
self-labelling (cell A2, e.g. 2026第1季合併財務報表) and carries the FX
volatility reserve (外匯價格變動準備) as a balance-sheet line through 2025;
the IFRS-17-era 2026 workbooks fold it into 其他負債, so the 2026 balance is
derived exactly as the 2025Q4 balance plus the cash-flow statement's YTD
外匯價格變動準備淨變動.

Checks per workbook: the balance-sheet current column's serial date matches the
label's quarter end; assets = liabilities + equity; and where the deck series
carries the reserve, statement and deck agree to deck rounding (the decisions
3.6 join). Emits data/cathay_statements_quarterly.csv and
out/stage3_cathay_stmt_YYYYMMDD.sql (idempotent).

Units: statements are NT$ thousand; stored as NT$ mn (three decimals).
Vintage: the statement's period-end date — publication dates are not
recoverable (the CMS re-stamps Last-Modified on re-uploads).
"""
import csv
import datetime as dt
import json
import re
import sys
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parent.parent
CACHE = ROOT / "cache" / "cathay_stmts"
INDEX = ROOT / "config" / "cathay_life_statements.json"
BASE = "https://www.cathayholdings.com"

Q_END = {1: "03-31", 2: "06-30", 3: "09-30", 4: "12-31"}
Q_START = {1: "01-01", 2: "04-01", 3: "07-01", 4: "10-01"}


def serial_to_date(n):
    if isinstance(n, str):
        m = re.search(r"(20\d\d)年(\d{1,2})月(\d{1,2})日", n)
        assert m, f"unparseable column date {n!r}"
        return dt.date(int(m.group(1)), int(m.group(2)), int(m.group(3)))
    if isinstance(n, dt.datetime):
        return n.date()
    return dt.date(1899, 12, 30) + dt.timedelta(days=int(n))


def parse_label(label):
    m = re.search(r"(20\d\d)\D*第(\d)季", label)
    assert m, f"unparseable label {label!r}"
    return int(m.group(1)), int(m.group(2))


def find_row(ws, item_text, exact=True):
    for row in ws.iter_rows(values_only=True):
        it = row[1] if len(row) > 1 else None
        if isinstance(it, str):
            it = it.strip()
            if (it == item_text) if exact else (item_text in it):
                return row
    return None


def first_num(row, start):
    for c in row[start:]:
        if isinstance(c, (int, float)):
            return c
    return None


def extract(fn):
    wb = openpyxl.load_workbook(fn, read_only=True, data_only=True)
    bs = wb["合併資產負債表"]
    rows = list(bs.iter_rows(values_only=True, max_row=8))
    label = next(c for r in rows[:3] for c in r if isinstance(c, str) and "財務報" in c)
    year, q = parse_label(label)
    hdr = next(r for r in rows if r[0] == "代號")
    cur_end = serial_to_date(hdr[2])
    want_end = dt.date.fromisoformat(f"{year}-{Q_END[q]}")
    assert cur_end == want_end, f"{fn.name}: current column {cur_end} != {want_end}"

    def bsline(txt, exact=True):
        r = find_row(bs, txt, exact)
        return None if r is None else r[2]

    out = {
        "year": year, "q": q, "label": label,
        "fx_reserve_th": bsline("外匯價格變動準備"),
        "total_assets_th": bsline("資產總計"),
        "liabilities_th": bsline("負債總計"),
        "equity_th": bsline("權益總計") if bsline("權益總計") is not None else bsline("權益總額"),
    }
    cf = wb["合併現金流量表"]
    r = find_row(cf, "外匯價格變動準備", exact=False)
    out["cf_fx_net_change_ytd_th"] = None if r is None else first_num(r, 2)
    wb.close()
    return out


def main():
    idx = json.loads(INDEX.read_text(encoding="utf-8"))
    import hashlib
    recs = []
    for i, e in enumerate(idx):
        link = e["links"].get("Excel")
        if not link:
            continue
        h = hashlib.md5(link.encode()).hexdigest()[:10]
        fn = CACHE / f"{i:02d}_{e['section']}_{h}.xlsx"
        assert fn.exists(), f"missing {fn} — download pass first"
        rec = extract(fn)
        rec["url"] = BASE + link
        recs.append(rec)

    recs.sort(key=lambda r: (r["year"], r["q"]))
    checks_total = checks_failed = 0

    # identity: assets = liabilities + equity
    for r in recs:
        checks_total += 1
        if r["total_assets_th"] != r["liabilities_th"] + r["equity_th"]:
            checks_failed += 1
            print(f"  FAIL identity A=L+E {r['year']}Q{r['q']}")

    # 2026 reserve derivation from the 2025Q4 balance + CF YTD net change
    base_row = next(r for r in recs if (r["year"], r["q"]) == (2025, 4))
    for r in recs:
        if r["fx_reserve_th"] is None:
            assert r["year"] >= 2026, f"{r['year']}Q{r['q']} missing reserve pre-IFRS17"
            assert r["cf_fx_net_change_ytd_th"] is not None
            r["fx_reserve_th"] = base_row["fx_reserve_th"] + r["cf_fx_net_change_ytd_th"]
            r["derived"] = True
        else:
            r["derived"] = False

    # deck join: where the merged deck series carries the reserve, agree to 0.05bn
    deck = {}
    with open(ROOT / "data" / "cathay_fx_quarterly.csv", encoding="utf-8") as f:
        for d in csv.DictReader(f):
            m = re.fullmatch(r"(FY|1Q|1H|9M)(\d{2})", d["period"] or "")
            if m and d["fx_volatility_reserve_ntd_bn"]:
                qq = {"1Q": 1, "1H": 2, "9M": 3, "FY": 4}[m.group(1)]
                deck[(2000 + int(m.group(2)), qq)] = float(d["fx_volatility_reserve_ntd_bn"])
    joined = 0
    for r in recs:
        want = deck.get((r["year"], r["q"]))
        if want is None:
            continue
        checks_total += 1
        joined += 1
        got_bn = r["fx_reserve_th"] / 1e6
        if abs(got_bn - want) > 0.05 + 1e-9:
            checks_failed += 1
            print(f"  FAIL deck join {r['year']}Q{r['q']}: stmt {got_bn:.3f}bn vs deck {want}bn")

    now = dt.datetime.now(dt.timezone.utc).isoformat()
    csv_path = ROOT / "data" / "cathay_statements_quarterly.csv"
    with open(csv_path, "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        w.writerow(["obs_quarter", "vintage", "basis", "fx_reserve_balance_mn",
                    "total_assets_mn", "owners_equity_mn", "reserve_derived",
                    "source_url", "source_doc"])
        for r in recs:
            r["obs_quarter"] = f"{r['year']}-{Q_START[r['q']]}"
            r["vintage"] = f"{r['year']}-{Q_END[r['q']]}"
            r["basis"] = "IFRS17" if r["year"] >= 2026 else "IFRS4"
            r["doc"] = f"國泰人壽 (5846) {r['label']} 合併資產負債表/現金流量表"
            w.writerow([r["obs_quarter"], r["vintage"], r["basis"],
                        round(r["fx_reserve_th"] / 1e3, 3),
                        round(r["total_assets_th"] / 1e3, 3),
                        round(r["equity_th"] / 1e3, 3),
                        r["derived"], r["url"], r["doc"]])

    def sqlv(v):
        if v is None:
            return "null"
        if isinstance(v, str):
            return "'" + v.replace("'", "''") + "'"
        return str(v)

    cols = "obs_quarter, vintage, basis, fx_res, ta, eq, derived, url, doc"
    vals = ",\n  ".join(
        "(" + ", ".join([
            sqlv(r["obs_quarter"]), sqlv(r["vintage"]), sqlv(r["basis"]),
            sqlv(round(r["fx_reserve_th"] / 1e3, 3)),
            sqlv(round(r["total_assets_th"] / 1e3, 3)),
            sqlv(round(r["equity_th"] / 1e3, 3)),
            "true" if r["derived"] else "false",
            sqlv(r["url"]), sqlv(r["doc"]),
        ]) + ")" for r in recs)
    sql = f"""-- stage3-cathay-stmt: generated by scripts/stage3_cathay_statements.py at {now}. Idempotent.
begin;
with vals({cols}) as (values
  {vals})
insert into tlfx.firm_quarterly
  (entity_id, obs_quarter, vintage, basis, source_channel, fx_reserve_balance,
   total_assets, owners_equity, source_url, source_doc, source_note, retrieved_at)
select 'cathay_life', obs_quarter::date, vintage::date, basis::tlfx.accounting_basis,
  'statement', fx_res, ta, eq, url, doc,
  case when derived
    then 'NT$ mn from NT$-thousand statements; reserve derived: 2025Q4 balance + cash-flow YTD 外匯價格變動準備淨變動 (IFRS-17 balance sheet folds the reserve into 其他負債); vintage = period end (publication date unrecoverable)'
    else 'NT$ mn from NT$-thousand statements; reserve is the balance-sheet line 外匯價格變動準備; vintage = period end (publication date unrecoverable)'
  end,
  '{now}'::timestamptz
from vals
on conflict (entity_id, obs_quarter, basis, source_channel, vintage) do nothing;
commit;
"""
    tag = dt.date.today().strftime("%Y%m%d")
    out = ROOT / "out" / f"stage3_cathay_stmt_{tag}.sql"
    out.write_text(sql, encoding="utf-8")

    n = len(recs)
    s_res = round(sum(r["fx_reserve_th"] for r in recs) / 1e3, 3)
    s_ta = round(sum(r["total_assets_th"] for r in recs) / 1e3, 3)
    s_eq = round(sum(r["equity_th"] for r in recs) / 1e3, 3)
    print(f"rows: {n}; {recs[0]['year']}Q{recs[0]['q']} -> {recs[-1]['year']}Q{recs[-1]['q']}; "
          f"checks {checks_total - checks_failed}/{checks_total} (deck joins: {joined})")
    print(f"checksums: sum fx_reserve {s_res}; sum total_assets {s_ta}; sum equity {s_eq}")
    print(f"csv: {csv_path.relative_to(ROOT)}\nsql: {out.relative_to(ROOT)}")
    return 1 if checks_failed else 0


if __name__ == "__main__":
    sys.exit(main())
